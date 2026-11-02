#!/usr/bin/env python3
"""Build an .xlsx from a JSON spec or a CSV with openpyxl: live formulas, a styled frozen header,
number formats, column widths, totals, named cells and charts, without writing any code.

Usage:
  python3 build_xlsx.py --spec spec.json --out /workspace/outputs/model.xlsx
  python3 build_xlsx.py --csv data.csv --out /workspace/outputs/data.xlsx [--decimal ,] [--sheet-name Data]
  python3 build_xlsx.py --example            # print an example spec

Spec: {"sheets": [SHEET, ...], "names": {"growth": "Inputs!$B$1"}} - or a bare list of rows (one
sheet), or one SHEET object. SHEET:
  {"name": "Plan",
   "rows": [["Month", "Revenue", "Cost", "Margin"],        # row 1 is the header unless "header": false
            ["Jan", 1000, 600, "=B2-C2"],                  # a string starting with "=" is a live formula
            ["Feb", "=B2*1.05", 630, "=B3-C3"]],           # numbers are numbers: write 5100, not "5100"
   "header": true,                 # style + freeze row 1 + autofilter (default true)
   "formats": {"B": "#,##0", "Margin": "#,##0", "D:E": "0.0%", "B2:B3": "0.0%"},   # column, header text, A:B range or cell range
   "widths": {"A": 14, "Notes": 40},                       # default: fitted to the content, 10-50
   "inputs": ["B2:B3"],            # blue font (inputs apart from calculations)
   "total": {"label": "Total", "cols": ["B", "C", "D"], "func": "SUM"},   # a formula row under the data
   "freeze": "A2",                 # default under the header; "none" to disable
   "charts": [{"type": "column", "title": "Revenue", "categories": "A2:A13", "series": ["B1:B13"],
               "anchor": "G2"}]}   # type: column | bar | line | pie; a series range starts with its header cell
Dates (datetime.date as ISO "2026-01-31") stay text unless the sheet or the CSV has "coerce": true /
--coerce, which turns numeric-looking text and ISO dates into numbers and dates.
Prints {"ok", "out", "sheets": {name: {"rows", "cols", "formulas"}}, "warnings", "next_step"}; a bad spec prints
every problem with the accepted keys (and the closest spelling) and exits 1. Then run recalc.py (the gate).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from opserr import OpError, did_you_mean  # noqa: E402

SHEET_KEYS = ["name", "rows", "header", "formats", "widths", "inputs", "total", "freeze", "charts", "coerce"]
CHART_KEYS = ["type", "title", "categories", "series", "anchor", "width_cm", "height_cm", "y_title", "x_title"]
CHART_TYPES = ["column", "bar", "line", "pie"]
# functions LibreOffice 7.4 (the recalculation check) does not compute: #NAME? in the gate
UNSUPPORTED = {"XLOOKUP", "XMATCH", "LET", "LAMBDA", "FILTER", "UNIQUE", "SORT", "SORTBY", "SEQUENCE",
               "RANDARRAY", "TEXTSPLIT", "VSTACK", "HSTACK", "TOCOL", "TOROW", "CHOOSECOLS", "CHOOSEROWS", "TAKE", "DROP"}
HEADER_FILL, HEADER_FONT, INPUT_COLOR = "1F3864", "FFFFFF", "2E6DA4"
EXAMPLE = {
    "names": {"growth": "Inputs!$B$1"},
    "sheets": [
        {"name": "Inputs", "header": False, "rows": [["Growth rate", 0.05], ["Tax rate", 0.25]],
         "formats": {"B": "0.0%"}, "inputs": ["B1:B2"], "widths": {"A": 18}},
        {"name": "Plan",
         "rows": [["Month", "Revenue", "Cost", "Margin", "Margin %"],
                  ["Jan", 1000, 600, "=B2-C2", "=IFERROR(D2/B2,0)"],
                  ["Feb", "=B2*(1+growth)", 630, "=B3-C3", "=IFERROR(D3/B3,0)"],
                  ["Mar", "=B3*(1+growth)", 660, "=B4-C4", "=IFERROR(D4/B4,0)"]],
         "formats": {"B:D": "#,##0", "Margin %": "0.0%"},
         "total": {"label": "Total", "cols": ["B", "C", "D"]},
         "charts": [{"type": "column", "title": "Revenue vs cost", "categories": "A2:A4",
                     "series": ["B1:B4", "C1:C4"], "anchor": "G2"}]},
    ],
}


def col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def unknown_keys(obj, allowed, where):
    problems = []
    for k in obj:
        if k not in allowed:
            hint = did_you_mean(k, allowed)
            problems.append(f"{where}: unknown key {k!r}" + (f" - did you mean {hint!r}?" if hint else "")
                            + f" (accepted: {', '.join(allowed)})")
    return problems


def normalize(spec):
    """The spec as {"sheets": [...], "names": {...}}; raises OpError listing every problem."""
    if isinstance(spec, list):
        spec = {"sheets": [{"name": "Sheet1", "rows": spec}]}
    elif isinstance(spec, dict) and "sheets" not in spec and "rows" in spec:
        spec = {"sheets": [spec]}
    problems = []
    if not isinstance(spec, dict) or not isinstance(spec.get("sheets"), list) or not spec["sheets"]:
        raise OpError("bad_spec", "the spec needs a non-empty \"sheets\" list (or a bare list of rows)",
                      usage="python3 build_xlsx.py --example")
    problems += unknown_keys(spec, ["sheets", "names"], "spec")
    seen = set()
    for i, sh in enumerate(spec["sheets"]):
        where = f"sheets[{i}]"
        if not isinstance(sh, dict):
            problems.append(f"{where}: a sheet is an object {{\"name\", \"rows\"}}")
            continue
        problems += unknown_keys(sh, SHEET_KEYS, where)
        name = str(sh.get("name") or f"Sheet{i + 1}")
        if len(name) > 31 or re.search(r"[\[\]:*?/\\]", name):
            problems.append(f"{where}: sheet name {name!r} must be at most 31 characters without []:*?/\\")
        if name.lower() in seen:
            problems.append(f"{where}: duplicate sheet name {name!r}")
        seen.add(name.lower())
        rows = sh.get("rows")
        if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
            problems.append(f"{where}: \"rows\" is a list of rows, each row a list of cells")
        for j, ch in enumerate(sh.get("charts") or []):
            problems += unknown_keys(ch, CHART_KEYS, f"{where}.charts[{j}]")
            if ch.get("type") not in CHART_TYPES:
                problems.append(f"{where}.charts[{j}]: type {ch.get('type')!r} is not one of {CHART_TYPES}")
            if not ch.get("series"):
                problems.append(f"{where}.charts[{j}]: \"series\" is a list of ranges such as [\"B1:B13\"]")
    if problems:
        raise OpError("bad_spec", f"{len(problems)} problem(s) in the spec", problems=problems,
                      usage="python3 build_xlsx.py --example  (prints a full spec)")
    spec["sheets"] = [dict(s, name=str(s.get("name") or f"Sheet{i + 1}")) for i, s in enumerate(spec["sheets"])]
    return spec


def coerce_value(v, decimal="."):
    """CSV/--coerce: numeric-looking text -> number, ISO date -> date."""
    if not isinstance(v, str):
        return v
    s = v.strip()
    if s == "":
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        try:
            return dt.date.fromisoformat(s)
        except ValueError:
            return v
    t = s.replace(" ", "").replace(" ", "")
    if decimal == ",":
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", t) else t
    if re.fullmatch(r"-?\d+", t):
        return int(t) if not (len(t.lstrip("-")) > 1 and t.lstrip("-").startswith("0")) else v
    if re.fullmatch(r"-?(\d+\.\d*|\.\d+)([eE]-?\d+)?|-?\d+[eE]-?\d+", t):
        return float(t)
    return v


def resolve_cols(key, headers):
    """'B', 'B:D' or a header text -> 1-based column indexes."""
    from openpyxl.utils import column_index_from_string
    key = str(key)
    m = re.fullmatch(r"([A-Za-z]{1,3})(?::([A-Za-z]{1,3}))?", key)
    if m and key not in headers:
        a = column_index_from_string(m.group(1).upper())
        b = column_index_from_string((m.group(2) or m.group(1)).upper())
        return list(range(a, b + 1))
    if key in headers:
        return [headers.index(key) + 1]
    hint = did_you_mean(key, headers)
    raise OpError("column_not_found", f"no column {key!r} (a letter, an A:C range or a header text)",
                  available=[str(h) for h in headers if h], did_you_mean=hint)


def build(spec, out, coerce=False, decimal="."):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import range_boundaries
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    wb.remove(wb.active)
    report, warnings = {}, []
    for sh in spec["sheets"]:
        ws = wb.create_sheet(sh["name"])
        do_coerce = coerce or sh.get("coerce")
        rows = [[coerce_value(c, decimal) if do_coerce else c for c in r] for r in sh["rows"]]
        header = sh.get("header", True) and bool(rows)
        formulas, text_numbers = 0, 0
        for r_i, row in enumerate(rows, 1):
            for c_i, v in enumerate(row, 1):
                if v is None:
                    continue
                cell = ws.cell(r_i, c_i, v)
                if isinstance(v, str) and v.startswith("="):
                    formulas += 1
                    for fn in re.findall(r"\b([A-Z][A-Z0-9.]+)\(", v.upper()):
                        if fn in UNSUPPORTED and f"unsupported {fn}" not in warnings:
                            warnings.append(f"unsupported {fn}: LibreOffice 7.4 gives #NAME? in the recalculation check")
                elif isinstance(v, str) and re.fullmatch(r"-?\d+([.,]\d+)?", v.strip()) and not (header and r_i == 1):
                    text_numbers += 1
                elif isinstance(v, dt.date) and not isinstance(v, dt.datetime):
                    cell.number_format = "yyyy-mm-dd"
        if text_numbers:
            warnings.append(f"{sh['name']}: {text_numbers} number(s) written as text (SUM counts text as 0): "
                            "write 5100 not \"5100\", or set \"coerce\": true")
        headers = [str(c) if c is not None else "" for c in (rows[0] if rows else [])] if header else []
        ncols = max((len(r) for r in rows), default=0)
        last = len(rows)

        if sh.get("total"):
            tot = sh["total"]
            func = str(tot.get("func", "SUM")).upper()
            r_t = last + 1
            ws.cell(r_t, 1, tot.get("label", "Total")).font = Font(bold=True)
            for key in tot.get("cols", []):
                for c in resolve_cols(key, headers):
                    L = col_letter(c)
                    cell = ws.cell(r_t, c, f"={func}({L}{2 if header else 1}:{L}{last})")
                    cell.font = Font(bold=True)
                    formulas += 1
            last = r_t
        for key, fmt in (sh.get("formats") or {}).items():
            if re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(:\$?[A-Za-z]{1,3}\$?\d+)?", str(key)):   # a cell range: "B2:B3", "C5"
                for row in ws[str(key)] if ":" in str(key) else [[ws[str(key)]]]:
                    for cell in row:
                        cell.number_format = fmt
                continue
            for c in resolve_cols(key, headers):
                for r in range(2 if header else 1, last + 1):
                    ws.cell(r, c).number_format = fmt
        for ref in sh.get("inputs") or []:
            for row in ws[ref] if ":" in ref else [[ws[ref]]]:
                for cell in row:
                    cell.font = Font(color=INPUT_COLOR)
        if header:
            for c in range(1, ncols + 1):
                cell = ws.cell(1, c)
                cell.font = Font(bold=True, color=HEADER_FONT)
                cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            if last >= 1 and ncols:
                ws.auto_filter.ref = f"A1:{col_letter(ncols)}{max(last - (1 if sh.get('total') else 0), 1)}"
        freeze = sh.get("freeze", "A2" if header else "none")
        if freeze and freeze != "none":
            ws.freeze_panes = freeze
        widths = {}
        for c in range(1, ncols + 1):
            longest = max((len(str(r[c - 1])) if isinstance(r[c - 1], str) and not r[c - 1].startswith("=") else 12
                           for r in rows if len(r) >= c and r[c - 1] is not None), default=10)
            widths[c] = min(max(longest + 2, 10), 50)
        for key, w in (sh.get("widths") or {}).items():
            for c in resolve_cols(key, headers):
                widths[c] = w
        for c, w in widths.items():
            ws.column_dimensions[col_letter(c)].width = w

        for ch in sh.get("charts") or []:
            kind = {"column": BarChart, "bar": BarChart, "line": LineChart, "pie": PieChart}[ch["type"]]
            chart = kind()
            if ch["type"] in ("column", "bar"):
                chart.type = "col" if ch["type"] == "column" else "bar"
            if ch.get("title"):
                chart.title = ch["title"]
            if ch.get("y_title") and hasattr(chart, "y_axis"):
                chart.y_axis.title = ch["y_title"]
            if ch.get("x_title") and hasattr(chart, "x_axis"):
                chart.x_axis.title = ch["x_title"]
            for rng in ch["series"]:
                c1, r1, c2, r2 = range_boundaries(rng)
                chart.add_data(Reference(ws, min_col=c1, max_col=c2, min_row=r1, max_row=r2), titles_from_data=True)
            if ch.get("categories"):
                c1, r1, c2, r2 = range_boundaries(ch["categories"])
                chart.set_categories(Reference(ws, min_col=c1, max_col=c2, min_row=r1, max_row=r2))
            chart.width, chart.height = ch.get("width_cm", 16), ch.get("height_cm", 8)
            ws.add_chart(chart, ch.get("anchor", f"{col_letter(ncols + 2)}2"))
        report[sh["name"]] = {"rows": len(rows), "cols": ncols, "formulas": formulas}
    for name, ref in (spec.get("names") or {}).items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    wb.save(out)
    return report, warnings


def read_csv(path, decimal, sheet_name):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        first = fh.readline()
        sep_line = first.strip().lower().startswith("sep=")
        delim = first.strip()[4:5] if sep_line else None
        if not sep_line:
            fh.seek(0)
        sample = fh.read(4096)
        fh.seek(0)
        if sep_line:
            fh.readline()
        if delim is None:
            try:
                delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delim = ","
        rows = [r for r in csv.reader(fh, delimiter=delim)]
    return {"sheets": [{"name": sheet_name, "rows": rows, "coerce": True}]}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", help="JSON spec (a file, or - for stdin)")
    ap.add_argument("--csv", help="build a one-sheet workbook from a CSV (a 'sep=;' first line is honoured)")
    ap.add_argument("--out", help="the .xlsx to write")
    ap.add_argument("--decimal", default=".", choices=[".", ","], help="decimal mark of the CSV (',' for 1.234,5)")
    ap.add_argument("--sheet-name", default="Data", help="sheet name for --csv")
    ap.add_argument("--coerce", action="store_true", help="numeric-looking text and ISO dates become numbers and dates")
    ap.add_argument("--example", action="store_true", help="print an example spec")
    args = ap.parse_args(argv)
    if args.example:
        print(json.dumps(EXAMPLE, indent=2, ensure_ascii=False))
        return 0
    if not args.out or not (args.spec or args.csv):
        ap.error("give --spec FILE (or --csv FILE) and --out BOOK.xlsx")
    if not args.out.lower().endswith((".xlsx", ".xlsm")):
        ap.error("--out must end in .xlsx")
    try:
        if args.csv:
            spec = read_csv(args.csv, args.decimal, args.sheet_name)
        else:
            spec = json.load(sys.stdin if args.spec == "-" else open(args.spec, encoding="utf-8"))
        report, warnings = build(normalize(spec), args.out, args.coerce, args.decimal)
    except OpError as e:
        print(json.dumps({"ok": False, **e.record()}, indent=2, ensure_ascii=False))
        return 1
    except (OSError, ValueError) as e:
        print(json.dumps({"ok": False, "reason": "unreadable_input", "error": f"{type(e).__name__}: {e}"}))
        return 1
    out = {"ok": True, "out": args.out, "sheets": report, "warnings": warnings,
           "next_step": f"python3 /skills/xlsx/scripts/recalc.py {args.out}  (computes the formulas, then the gate)"}
    print(json.dumps(out, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - bad input: one JSON error on stdout, no traceback
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
