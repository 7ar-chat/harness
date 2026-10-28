#!/usr/bin/env python3
"""Set cell values or formulas directly in the sheet XML of an .xlsx/.xlsm, preserving
everything else (styles, charts, pivots, slicers, macros, shapes). Use it to edit workbooks
that openpyxl would damage (see inspect_xlsx.py "fragile_parts"), then run recalc.py.

Usage:
  python3 set_cells.py BOOK.xlsx --out NEW.xlsx --set "Data!B2=42" --set "Data!C9==SUM(C2:C8)"
  python3 set_cells.py BOOK.xlsx --inplace --set "Summary!A1=Q3 2026 review"
  python3 set_cells.py BOOK.xlsx --inplace --json edits.json
      edits.json: [{"sheet": "Data", "cell": "B2", "value": 42},
                   {"sheet": "Data", "cell": "C9", "formula": "SUM(C2:C8)"},
                   {"sheet": "Data", "cell": "A10", "value": "Total", "style_from": "A9"}]

--set "Sheet!REF=VALUE": VALUE starting with "=" is a formula; numbers become numbers;
everything else is text. Use --json for explicit types (value null clears the cell).
Prints {"out", "changed": n, "sheets": {...}}. Run recalc.py afterwards to compute formulas.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

from lxml import etree

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def q(tag):
    return "{%s}%s" % (S, tag)


def split_ref(ref):
    m = re.match(r"^\$?([A-Za-z]{1,3})\$?(\d+)$", ref)
    if not m:
        raise SystemExit(f"error: bad cell reference {ref!r}")
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(m.group(1).upper()), int(m.group(2)), m.group(1).upper() + m.group(2)


def sheet_parts(zf):
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for sh in wb.find(q("sheets")):
        target = targets.get(sh.get("{%s}id" % R), "")
        out[sh.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    return out


def parse_set(text):
    m = re.match(r"^(.+?)!(\$?[A-Za-z]{1,3}\$?\d+)=(.*)$", text, re.S)
    if not m:
        raise SystemExit(f"error: --set expects 'Sheet!A1=value', got {text!r}")
    sheet, cell, raw = m.group(1).strip("'"), m.group(2), m.group(3)
    edit = {"sheet": sheet, "cell": cell}
    if raw.startswith("="):
        edit["formula"] = raw[1:]
    else:
        try:
            edit["value"] = int(raw) if re.match(r"^-?\d+$", raw) else float(raw) if re.match(
                r"^-?\d*\.\d+(e-?\d+)?$|^-?\d+e-?\d+$", raw, re.I) else raw
        except ValueError:
            edit["value"] = raw
    return edit


def get_row(data, r):
    for row in data.findall(q("row")):
        n = int(row.get("r"))
        if n == r:
            return row
        if n > r:
            new = etree.Element(q("row"), {"r": str(r)})
            row.addprevious(new)
            return new
    return etree.SubElement(data, q("row"), {"r": str(r)})


def get_cell(row, col, ref):
    from openpyxl.utils import column_index_from_string

    for c in row.findall(q("c")):
        cref = c.get("r")
        ccol = column_index_from_string(re.match(r"[A-Z]+", cref).group(0))
        if ccol == col:
            return c
        if ccol > col:
            new = etree.Element(q("c"), {"r": ref})
            c.addprevious(new)
            return new
    return etree.SubElement(row, q("c"), {"r": ref})


def expand_shared(data, si):
    """Replace a shared-formula group by explicit formulas (so its master can be edited)."""
    from openpyxl.formula.translate import Translator

    master = None
    for c in data.iter(q("c")):
        f = c.find(q("f"))
        if f is not None and f.get("t") == "shared" and f.get("si") == si and f.text:
            master = (c.get("r"), f.text)
            break
    if master is None:
        return
    for c in data.iter(q("c")):
        f = c.find(q("f"))
        if f is not None and f.get("t") == "shared" and f.get("si") == si:
            text = f.text if c.get("r") == master[0] else Translator("=" + master[1], origin=master[0]).translate_formula(c.get("r"))[1:]
            for attr in ("t", "si", "ref"):
                f.attrib.pop(attr, None)
            f.text = text


def apply(c, edit, data):
    f = c.find(q("f"))
    if f is not None and f.get("t") == "shared" and f.get("ref"):
        expand_shared(data, f.get("si"))
    for child in list(c):
        if child.tag in (q("f"), q("v"), q("is")):
            c.remove(child)
    c.attrib.pop("t", None)
    if "formula" in edit:
        f = etree.SubElement(c, q("f"))
        f.text = edit["formula"].lstrip("=")
        return True
    value = edit.get("value")
    if value is None:
        return False
    if isinstance(value, bool):
        c.set("t", "b")
        etree.SubElement(c, q("v")).text = "1" if value else "0"
    elif isinstance(value, (int, float)):
        etree.SubElement(c, q("v")).text = repr(value)
    else:
        c.set("t", "inlineStr")
        is_ = etree.SubElement(c, q("is"))
        t = etree.SubElement(is_, q("t"))
        t.text = str(value)
        t.set(XML_SPACE, "preserve")
    return False


def update_dimension(root, data):
    from openpyxl.utils import get_column_letter, column_index_from_string

    cols, rows = [], []
    for c in data.iter(q("c")):
        m = re.match(r"([A-Z]+)(\d+)", c.get("r"))
        cols.append(column_index_from_string(m.group(1)))
        rows.append(int(m.group(2)))
    dim = root.find(q("dimension"))
    if dim is not None and cols:
        dim.set("ref", f"{get_column_letter(min(cols))}{min(rows)}:{get_column_letter(max(cols))}{max(rows)}")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--set", action="append", default=[], dest="sets")
    ap.add_argument("--json")
    ap.add_argument("--out")
    ap.add_argument("--inplace", action="store_true")
    args = ap.parse_args(argv)
    if not (args.out or args.inplace):
        ap.error("give --out NEW.xlsx or --inplace")
    edits = [parse_set(s) for s in args.sets]
    if args.json:
        edits += json.load(open(args.json, encoding="utf-8"))
    if not edits:
        ap.error("no edits (use --set or --json)")

    zin = zipfile.ZipFile(args.file)
    parts = sheet_parts(zin)
    by_part = {}
    for e in edits:
        if e["sheet"] not in parts:
            raise SystemExit(f"error: no sheet {e['sheet']!r}; sheets: {list(parts)}")
        by_part.setdefault(parts[e["sheet"]], []).append(e)
    new_parts, formula_touched, per_sheet = {}, False, {}
    for part, part_edits in by_part.items():
        root = etree.fromstring(zin.read(part))
        data = root.find(q("sheetData"))
        for e in part_edits:
            col, row_n, ref = split_ref(e["cell"])
            row = get_row(data, row_n)
            c = get_cell(row, col, ref)
            had_formula = c.find(q("f")) is not None
            if e.get("style_from"):
                scol, srow, sref = split_ref(e["style_from"])
                src = data.find(f".//{q('c')}[@r='{sref}']")
                if src is not None and src.get("s"):
                    c.set("s", src.get("s"))
            is_formula = apply(c, e, data)
            formula_touched |= is_formula or had_formula
            per_sheet[e["sheet"]] = per_sheet.get(e["sheet"], 0) + 1
        update_dimension(root, data)
        new_parts[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    drop = set()
    if formula_touched and "xl/calcChain.xml" in zin.namelist():
        drop.add("xl/calcChain.xml")
        rels = etree.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        for r in list(rels):
            if r.get("Target", "").endswith("calcChain.xml"):
                rels.remove(r)
        new_parts["xl/_rels/workbook.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        ct = etree.fromstring(zin.read("[Content_Types].xml"))
        for o in list(ct):
            if o.get("PartName", "").endswith("calcChain.xml"):
                ct.remove(o)
        new_parts["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)

    out = args.file if args.inplace else args.out
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(out)) or ".")
    os.close(fd)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in drop:
                continue
            zout.writestr(item, new_parts.get(item.filename, zin.read(item.filename)))
    zin.close()
    shutil.move(tmp, out)
    print(json.dumps({"out": out, "changed": len(edits), "sheets": per_sheet,
                      "next": "run recalc.py to compute formula results" if formula_touched else None}, indent=2))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - bad input: one JSON error on stdout, no traceback
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
