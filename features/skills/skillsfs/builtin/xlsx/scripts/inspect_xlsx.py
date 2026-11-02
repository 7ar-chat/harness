#!/usr/bin/env python3
"""Summarise an .xlsx/.xlsm, or dump a cell range.

Usage:
  python3 inspect_xlsx.py BOOK.xlsx                         # JSON workbook summary
  python3 inspect_xlsx.py BOOK.xlsx --sheet Data --range A1:H30            # formulas as written (= --formulas)
  python3 inspect_xlsx.py BOOK.xlsx --sheet Data --range A1:H30 --values   # cached values (run recalc.py first)
  python3 inspect_xlsx.py BOOK.xlsx --formulas                            # every formula of every sheet (--sheet: one)
  python3 inspect_xlsx.py BOOK.xlsx --sheet Data --stats                  # used range, counts per kind, header: no cells
  python3 inspect_xlsx.py BOOK.xlsx --sheet Data --where error --values   # only the cells of one kind: formula|error|empty|number|text
  python3 inspect_xlsx.py BOOK.xlsx --sheet Data --range A1:H900 --cols A,C:E --max-rows 50   # bounded dump

A dump shows at most --max-rows rows (default 100) and cell text is cut at --max-chars (default 200) with
the marker …(+n chars); a cut dump says which --range continues it ("note"). --where lists {ref, value, formula}
only, up to 200 matches. A wrong --sheet answers with the sheet names and the closest spelling.

Summary per sheet: dimensions, header row, freeze panes, autofilter, tables, merged cells,
formula count + samples, conditional formats, data validations, number formats, column
widths. Workbook: defined names, and a scan of the package for charts, images, pivot
tables, slicers, macros, external links and comments. openpyxl re-writes or loses some of
these when it re-saves a file, so check "fragile_parts" before editing with openpyxl.
"""
from __future__ import annotations

import argparse
import collections
import json
import re
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from opserr import OpError, clip, did_you_mean  # noqa: E402


def scan_package(path):
    names = zipfile.ZipFile(path).namelist()
    pat = {"charts": r"xl/charts/chart\d+\.xml$", "images": r"xl/media/", "drawings": r"xl/drawings/drawing\d+\.xml$",
           "pivot_tables": r"xl/pivotTables/", "pivot_caches": r"xl/pivotCache/", "slicers": r"xl/slicers?/",
           "vba_macros": r"xl/vbaProject\.bin$", "external_links": r"xl/externalLinks/", "comments": r"xl/comments\d*\.xml$",
           "threaded_comments": r"xl/threadedComments/", "shapes_vml": r"xl/drawings/vmlDrawing", "calc_chain": r"xl/calcChain\.xml$"}
    found = {k: sum(1 for n in names if re.search(v, n)) for k, v in pat.items()}
    fragile = [k for k in ("charts", "images", "pivot_tables", "slicers", "vba_macros", "external_links",
                           "threaded_comments") if found[k]]
    return found, fragile


def jsonable(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def pick_sheet(wb, sheet):
    if not sheet:
        return wb.active
    if sheet not in wb.sheetnames:
        raise OpError("sheet_not_found", f"sheet not found: {sheet!r}", available=wb.sheetnames,
                      did_you_mean=did_you_mean(sheet, wb.sheetnames))
    return wb[sheet]


def parse_cols(spec):
    from openpyxl.utils import column_index_from_string
    out = []
    for part in spec.split(","):
        a, _, b = part.strip().partition(":")
        out += list(range(column_index_from_string(a.upper()), column_index_from_string((b or a).upper()) + 1))
    return out


def cell_kind(c):
    v = c.value
    if v is None:
        return "empty"
    if isinstance(v, str) and v.startswith("#") and v.rstrip("!?").upper() in (
            "#REF", "#DIV/0", "#VALUE", "#NAME", "#N/A", "#NUM", "#NULL"):
        return "error"
    if c.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
        return "formula"
    return "number" if isinstance(v, (int, float)) and not isinstance(v, bool) else "text"


def dump_range(path, sheet, ref, values, max_rows=100, max_chars=200, cols=None):
    import openpyxl
    from openpyxl.utils import get_column_letter, range_boundaries

    wb = openpyxl.load_workbook(path, data_only=values)
    ws = pick_sheet(wb, sheet)
    c1, r1, c2, r2 = range_boundaries(ref)
    total = r2 - r1 + 1
    r_end = min(r2, r1 + max_rows - 1) if max_rows and max_rows > 0 else r2
    keep = [c for c in (cols or range(c1, c2 + 1)) if c1 <= c <= c2]
    rows = [[clip(jsonable(ws.cell(r, c).value), max_chars) if isinstance(ws.cell(r, c).value, str)
             else jsonable(ws.cell(r, c).value) for c in keep] for r in range(r1, r_end + 1)]
    out = {"sheet": ws.title, "range": f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r_end}",
           "values" if values else "cells": rows}
    if cols:
        out["columns"] = [get_column_letter(c) for c in keep]
    if r_end < r2:
        out.update(rowsShown=r_end - r1 + 1, rowsTotal=total,
                   note=f"{r_end - r1 + 1} of {total} rows shown; continue with --range "
                        f"{get_column_letter(c1)}{r_end + 1}:{get_column_letter(c2)}{min(r2, r_end + max_rows)}")
    return out


def where_cells(path, sheet, ref, kind, values, limit=200, max_chars=200):
    """Only the cells of one kind (formula|error|empty|number|text): [{ref, value, formula?}]."""
    import openpyxl

    wb = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True) if values or kind == "error" else None
    ws = pick_sheet(wb, sheet)
    wsv = wv[ws.title] if wv else None
    hits, total = [], 0
    for row in ws[ref] if ref else ws.iter_rows():
        for c in row:
            shown = wsv[c.coordinate].value if wsv is not None else c.value
            probe = c if wsv is None else wsv[c.coordinate]
            k = cell_kind(probe) if kind == "error" else cell_kind(c)
            if k != kind:
                continue
            total += 1
            if len(hits) < limit:
                item = {"ref": c.coordinate}
                if kind == "formula" and wsv is None:
                    item["formula"] = clip(str(c.value), max_chars)  # no cached value asked: the formula is the content
                else:
                    item["value"] = clip(jsonable(shown), max_chars)
                    if cell_kind(c) == "formula":
                        item["formula"] = clip(str(c.value), max_chars)
                hits.append(item)
    out = {"sheet": ws.title, "where": kind, "cells": hits, "matches": total}
    if total > len(hits):
        out["note"] = f"{len(hits)} of {total} matching cells listed; narrow --range"
    return out


def sheet_stats(path, sheet):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = pick_sheet(wb, sheet)
    kinds = collections.Counter(cell_kind(c) for row in ws.iter_rows() for c in row if c.value is not None)
    header = [clip(jsonable(c.value), 40) for c in next(ws.iter_rows(min_row=1, max_row=1), [])][:30]
    return {"sheet": ws.title, "sheets": wb.sheetnames, "used_range": ws.dimensions, "rows": ws.max_row,
            "columns": ws.max_column, "cells": dict(kinds), "header_row": header,
            "freeze_panes": ws.freeze_panes, "autofilter": ws.auto_filter.ref}


def list_formulas(path, sheet=None, limit=2000):
    import openpyxl

    wb = openpyxl.load_workbook(path)  # formulas as written
    sheets = [wb[sheet]] if sheet else wb.worksheets
    cells, total = [], 0
    for ws in sheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    total += 1
                    if len(cells) < limit:
                        cells.append({"sheet": ws.title, "cell": c.coordinate, "formula": str(c.value)})
    return {"file": path, "formulas": cells, "count": total, "truncated": total > len(cells)}


def summary(path):
    import openpyxl

    wb = openpyxl.load_workbook(path)  # formulas as written
    found, fragile = scan_package(path)
    sheets = []
    for ws in wb.worksheets:
        formulas = []
        fcount = 0
        nfmt = collections.Counter()
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    fcount += 1
                    if len(formulas) < 10:
                        formulas.append({"cell": c.coordinate, "formula": str(c.value)})
                if c.value is not None and c.number_format != "General":
                    nfmt[c.number_format] += 1
        header = [jsonable(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1), [])][:30]
        sheets.append({
            "name": ws.title, "state": ws.sheet_state, "dimensions": ws.dimensions,
            "max_row": ws.max_row, "max_column": ws.max_column, "header_row": header,
            "freeze_panes": ws.freeze_panes, "autofilter": ws.auto_filter.ref,
            "tables": [{"name": t.name, "ref": t.ref} for t in ws.tables.values()],
            "merged_ranges": len(ws.merged_cells.ranges),
            "formulas": fcount, "formula_samples": formulas,
            "conditional_formats": sum(len(r.rules) for r in ws.conditional_formatting),
            "data_validations": len(ws.data_validations.dataValidation),
            "number_formats": dict(nfmt.most_common(8)),
            "column_widths": {k: round(v.width, 1) for k, v in list(ws.column_dimensions.items())[:30] if v.width},
        })
    names = {}
    try:
        for name, dn in wb.defined_names.items():
            names[name] = dn.attr_text
    except AttributeError:  # older openpyxl
        names = {dn.name: dn.attr_text for dn in wb.defined_names.definedName}
    warnings = []
    if fragile:
        warnings.append("package contains " + ", ".join(fragile) + ": openpyxl load+save re-writes charts/images "
                        "(formatting can change) and loses shapes, slicers, form controls, threaded comments and "
                        "macros (unless keep_vba=True). For value/formula edits use set_cells.py, then recalc.py")
    return {"file": path, "sheets": sheets, "defined_names": names, "package": found,
            "fragile_parts": fragile, "warnings": warnings}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--range")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--values", action="store_true", help="cached values instead of formulas (run recalc.py first)")
    mode.add_argument("--formulas", action="store_true",
                      help="formulas as written (the default with --range); without --range: every formula cell of "
                           "the workbook (or of --sheet), up to 2000")
    ap.add_argument("--stats", action="store_true", help="with --sheet: used range, counts per cell kind, header; no cells")
    ap.add_argument("--where", choices=["formula", "error", "empty", "number", "text"],
                    help="list only the cells of this kind (with --values: cached values; error needs recalc.py first)")
    ap.add_argument("--cols", help="columns to keep in a --range dump, e.g. A,C:E")
    ap.add_argument("--max-rows", type=int, default=100, help="rows shown by a --range dump (default 100; 0 = all)")
    ap.add_argument("--max-chars", type=int, default=200, help="cell text cut at this many characters (default 200)")
    args = ap.parse_args(argv)
    try:
        if args.where:
            out = where_cells(args.file, args.sheet, args.range, args.where, args.values, max_chars=args.max_chars)
        elif args.stats:
            out = sheet_stats(args.file, args.sheet)
        elif args.range:
            out = dump_range(args.file, args.sheet, args.range, args.values, args.max_rows, args.max_chars,
                             parse_cols(args.cols) if args.cols else None)
        elif args.formulas:
            out = list_formulas(args.file, args.sheet)
        elif args.values:
            raise ValueError("--values needs --range (e.g. --sheet Summary --range A1:F15 --values)")
        else:
            out = summary(args.file)
    except OpError as exc:
        print(json.dumps({"ok": False, **exc.record()}, ensure_ascii=False))
        return 1
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"error": f"{type(exc).__name__}: {exc}"}))
        return 1
    print(json.dumps(out, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - bad input: one JSON error on stdout, no traceback
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
