"""Generate the golden fixtures of the Go classification port.

Sources are built with the same libraries Studio uses (python-docx, openpyxl,
python-pptx, reportlab) and labelled by the reference implementation
(ai/file_classification) with a fixed context, so the Go port can prove
property-for-property parity.

Run from the repository root (needs the ai image dependencies):

    docker run --rm -v "$PWD:/src" -w /src --entrypoint python \
        <studio-api-image>:<tag> \
        work/internal/classification/testdata/gen/gen_fixtures.py

Outputs (work/internal/classification/testdata/):
    <name>.<ext>            source document
    <name>.labelled.<ext>   output of ClassificationService.classify
    expected.json           context + properties read back from every output
"""

from __future__ import annotations

import io
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[5]
sys.path.insert(0, str(ROOT / "ai"))

import docx  # noqa: E402
import openpyxl  # noqa: E402
import pptx  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

from file_classification.configuration import load_classification_settings  # noqa: E402
from file_classification.ooxml import read_custom_properties, write_custom_properties  # noqa: E402
from file_classification.service import ClassificationService  # noqa: E402

OUT = Path(__file__).resolve().parents[1]

# Same defaults as ai/api/settings.py (FILE_CLASSIFICATION_*).
SETTINGS = {
    "mode": "enforce",
    "profile": {
        "name": "c3",
        "version": "2026-08",
        "level": "C3",
        "providers": ("mip", "gvdata"),
        "provider_settings": {
            "mip": {
                "enabled": "true",
                "label_id": "18dda7a4-b650-4246-b8ce-bc9857d399d0",
                "label_name": "C3-External-Allowed",
                "site_id": "a5c34232-eadc-4609-bff3-dd6fcdae3fe2",
            },
            "gvdata": {
                "enabled": "true",
                "classification": "Confidentiel",
                "classification_tag_set_id": "e16409a7-1700-4153-9090-3955bc2f0ae8",
                "compliance_tag_set_id": "f14fc1f1-8950-40d5-8a29-45909da947d6",
                "distribution": "CA-TITRES",
                "distribution_tag_set_id": "004dea33-8751-4399-a76e-95f371cb4119",
            },
        },
    },
    "actor": "work",
}

NOW = datetime(2026, 10, 5, 12, 34, 56, tzinfo=timezone.utc)
UUIDS = (
    "11111111-1111-4111-8111-111111111111",  # document_id
    "22222222-2222-4222-8222-222222222222",  # action_id
    "33333333-3333-4333-8333-333333333333",  # line_id
)


def service() -> ClassificationService:
    ids = iter(uuid.UUID(u) for u in UUIDS)
    loaded = load_classification_settings(SETTINGS)
    return ClassificationService(
        settings={**SETTINGS, "profile": loaded.profile},
        now=lambda: NOW,
        uuid_factory=lambda: next(ids),
    )


def make_docx() -> bytes:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Quarterly header"
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("Revenue grew by 12 % — écarts expliqués.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_docx_with_custom() -> bytes:
    # A document that already has custom properties (merge path).
    return write_custom_properties(make_docx(), {"Owner": "Finance", "Project": "Q3"})


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount"])
    ws.append(["EMEA", 1200])
    ws.append(["APAC", 900])
    ws["C2"] = "=B2*2"
    wb.create_sheet("Notes")["A1"] = "Draft"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Strategy 2027"
    slide.placeholders[1].text = "Grow, simplify, deliver"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595.27, 841.89))
    c.setTitle("Fixture")
    c.setAuthor("Studio")
    c.drawString(72, 770, "Hello classified world")
    c.showPage()
    c.drawString(72, 770, "Second page")
    c.showPage()
    c.save()
    return buf.getvalue()


FIXTURES = (
    ("report", "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx),
    ("custom", "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", make_docx_with_custom),
    ("sales", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", make_xlsx),
    ("deck", "pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", make_pptx),
    ("memo", "pdf", "application/pdf", make_pdf),
)


def properties_of(ext: str, content: bytes) -> dict[str, str]:
    if ext == "pdf":
        meta = PdfReader(io.BytesIO(content)).metadata or {}
        return {str(k): str(v) for k, v in meta.items()}
    return read_custom_properties(content)


def main() -> None:
    expected: dict[str, object] = {
        "timestamp": NOW.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "documentId": UUIDS[0],
        "actionId": UUIDS[1],
        "lineId": UUIDS[2],
        "actor": SETTINGS["actor"],
        "files": {},
    }
    for name, ext, mime, build in FIXTURES:
        source = build()
        (OUT / f"{name}.{ext}").write_bytes(source)
        result = service().classify(filename=f"{name}.{ext}", content_type=mime, content=source)
        (OUT / f"{name}.labelled.{ext}").write_bytes(result.content)
        expected["files"][f"{name}.{ext}"] = {
            "mimeType": mime,
            "sourceProperties": properties_of(ext, source),
            "properties": properties_of(ext, result.content),
        }
    (OUT / "expected.json").write_text(json.dumps(expected, indent=2, ensure_ascii=False, sort_keys=True) + "\n")
    print("wrote fixtures to", OUT)


if __name__ == "__main__":
    main()
