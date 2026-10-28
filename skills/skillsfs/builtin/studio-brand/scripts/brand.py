#!/usr/bin/env python3
"""Studio neutral corporate brand: constants + helpers for python-docx, python-pptx,
openpyxl / xlsxwriter and matplotlib.

Import it from any script:

    import sys; sys.path.insert(0, "/skills/studio-brand/scripts")
    import brand

CLI:
    python3 brand.py --json            # print palette / fonts / sizes as JSON
    python3 brand.py --swatch out.png  # render a palette swatch image
    python3 brand.py --format 1234567.891 --locale fr --decimals 2

Every helper imports its document library lazily, so importing this module only needs
the standard library.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import sys

# --------------------------------------------------------------------------- palette
# Sober, professional, colour-blind-aware. Hex without "#" is what python-docx /
# python-pptx / openpyxl expect; use hex_() for "#RRGGBB" (matplotlib, HTML, xlsxwriter).
PALETTE = {
    "navy": "1B365D",       # primary: titles, header rows, key series
    "blue": "2E6DA4",       # secondary
    "sky": "8DB8DE",        # light accent, backgrounds of highlights
    "teal": "2A9D8F",       # accent
    "amber": "E9A23B",      # highlight / warning
    "coral": "C8553D",      # alert / negative emphasis
    "ink": "1A1F2B",        # body text
    "slate": "4A5568",      # secondary text, captions
    "grey": "8C96A3",       # axis, footers, muted series
    "rule": "D5DAE1",       # table borders, separators, gridlines
    "mist": "EEF1F5",       # zebra rows, panels
    "white": "FFFFFF",
    "positive": "2E7D32",   # gains (use with a sign or icon, never colour alone)
    "negative": "C62828",   # losses
}

# Ordered categorical series colours (charts, legends). Use in this order; max 6-8 series.
CHART_COLORS = ["1B365D", "2E6DA4", "2A9D8F", "E9A23B", "C8553D", "7D5BA6", "8C96A3", "8AA05B"]
# Sequential (light -> dark) and diverging (negative -> neutral -> positive) ramps.
SEQUENTIAL = ["EEF1F5", "C9D8EA", "8DB8DE", "4F8BC0", "2E6DA4", "1B365D"]
DIVERGING = ["C62828", "E58A7B", "F2F2F2", "7FB3D5", "1B365D"]

# --------------------------------------------------------------------------- fonts
# Office files name "Arial" (installed on every Windows/macOS machine). In the sandbox
# LibreOffice renders it with the metric-compatible "Liberation Sans", so line breaks and
# page counts are the same as in Word/PowerPoint. Matplotlib must use an installed font.
FONTS = {
    "office": "Arial",                  # headings + body in docx / pptx / xlsx
    "office_mono": "Courier New",       # code / fixed width (rendered as Liberation Mono)
    "render": "Liberation Sans",        # matplotlib, Pillow, reportlab TTF
    "render_fallback": ["Liberation Sans", "DejaVu Sans", "Noto Sans"],
    "reportlab": "StudioSans",          # Liberation Sans TTF registered by reportlab_fonts()
}
FONT_FILES = {
    "regular": "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "bold": "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "italic": "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
    "bold_italic": "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf",
}

# Type scale in points.
DOCX_SIZES = {"title": 26, "subtitle": 14, "h1": 16, "h2": 13, "h3": 11.5, "body": 10.5,
              "small": 9, "caption": 9, "footer": 8, "table": 9.5}
PPTX_SIZES = {"title_slide": 40, "title": 28, "subtitle": 18, "body": 18, "body_min": 12,
              "small": 12, "footer": 10, "kpi": 44}
XLSX_SIZES = {"title": 14, "header": 10, "body": 10}

# Page geometry.
DOCX_PAGE = {"size": "A4", "width_cm": 21.0, "height_cm": 29.7, "margin_cm": 2.2,
             "text_width_cm": 16.6}
PPTX_SLIDE = {"width_in": 13.333, "height_in": 7.5, "margin_in": 0.6}

# Chart figure sizes (inches) matching the targets above; export at 200 dpi.
FIGURE_SIZES = {
    "slide": (11.5, 5.2),        # full-width content area of a 16:9 slide
    "slide-half": (5.8, 4.8),    # one of two columns on a slide
    "doc": (6.5, 3.6),           # full text width of an A4 portrait page
    "doc-half": (3.2, 2.6),
    "square": (5.0, 5.0),
}
DPI = 200


# Colours in code: each library wants its own form (a wrong one raises "Colors must be aRGB hex
# values" in openpyxl or "rgb color value must be RGBColor object" in python-docx / python-pptx).
#   openpyxl     Font(color=xlsx_color("navy")), PatternFill("solid", fgColor=xlsx_color("mist")),
#                cell.fill = xlsx_fill("mist"), cell.font = xlsx_font("white", bold=True)
#   python-docx  run.font.color.rgb = docx_rgb("slate")
#   python-pptx  shape.fill.fore_color.rgb = pptx_rgb("navy")
#   matplotlib / xlsxwriter / reportlab / HTML: hex_("navy") -> "#1B365D";  Pillow: rgb("navy")
def _hex6(name_or_hex: str) -> str:
    """'navy', 'Navy', '1b365d' or '#1B365D' -> '1B365D' (ValueError for anything else)."""
    text = str(name_or_hex).strip()
    value = PALETTE.get(text.lower(), text).lstrip("#").upper()
    if len(value) != 6 or any(ch not in "0123456789ABCDEF" for ch in value):
        raise ValueError(f"unknown colour {name_or_hex!r}: use a palette name ({', '.join(PALETTE)}) "
                         f"or a hex code 'RRGGBB'")
    return value


def hex_(name_or_hex: str) -> str:
    """'navy' or '1B365D' -> '#1B365D': for matplotlib, xlsxwriter, reportlab and HTML only.
    openpyxl rejects the '#': use xlsx_color()."""
    return "#" + _hex6(name_or_hex)


def rgb(name_or_hex: str) -> tuple[int, int, int]:
    """'navy' -> (27, 54, 93): a plain tuple for Pillow. python-docx / python-pptx need
    docx_rgb() / pptx_rgb()."""
    value = _hex6(name_or_hex)
    return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)


def xlsx_color(name_or_hex: str) -> str:
    """'navy' -> 'FF1B365D' (aRGB): the colour argument of openpyxl's Font, PatternFill, Side,
    conditional-formatting rules and chart colours."""
    return "FF" + _hex6(name_or_hex)


def xlsx_fill(name_or_hex: str):
    """openpyxl solid fill: cell.fill = brand.xlsx_fill("mist")."""
    from openpyxl.styles import PatternFill

    return PatternFill("solid", fgColor=xlsx_color(name_or_hex))


def xlsx_font(color: str = "ink", bold: bool = False, italic: bool = False, size: float | None = None):
    """openpyxl Arial font in a palette colour: cell.font = brand.xlsx_font("white", bold=True)."""
    from openpyxl.styles import Font

    return Font(name=FONTS["office"], size=size or XLSX_SIZES["body"], bold=bold, italic=italic,
                color=xlsx_color(color))


def docx_rgb(name_or_hex: str):
    """python-docx colour: run.font.color.rgb = brand.docx_rgb("slate")."""
    from docx.shared import RGBColor

    return RGBColor.from_string(_hex6(name_or_hex))


def pptx_rgb(name_or_hex: str):
    """python-pptx colour: shape.fill.fore_color.rgb = brand.pptx_rgb("navy")."""
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(_hex6(name_or_hex))


# --------------------------------------------------------------------------- formatting
MONTHS = {
    "fr": ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
           "septembre", "octobre", "novembre", "décembre"],
    "en": ["January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December"],
}
NBSP = "\u00a0"         # before %, € and : in French
NNBSP = "\u202f"        # typographic narrow no-break space: NOT used, Liberation/Arial lack the glyph
FR_THOUSANDS = NBSP     # French thousands separator (same as Excel/Windows fr-FR)


def _norm_locale(locale: str) -> str:
    loc = (locale or "fr").lower()[:2]
    if loc not in ("fr", "en"):
        raise ValueError(f"unsupported locale {locale!r} (use 'fr' or 'en')")
    return loc


def format_number(value: float, decimals: int = 0, locale: str = "fr", sign: bool = False) -> str:
    """1234567.891 -> '1 234 567,89' (fr, no-break space) / '1,234,567.89' (en)."""
    loc = _norm_locale(locale)
    text = f"{abs(value):,.{decimals}f}"
    if loc == "fr":
        text = text.replace(",", FR_THOUSANDS).replace(".", ",")
    if value < 0 and float(f"{abs(value):.{decimals}f}") != 0:
        text = "-" + text
    elif sign and value > 0:
        text = "+" + text
    return text


def format_percent(value: float, decimals: int = 1, locale: str = "fr", ratio: bool = True,
                   sign: bool = False) -> str:
    """0.125 -> '12,5 %' (fr) / '12.5%' (en). ratio=False if value is already in percent."""
    loc = _norm_locale(locale)
    number = format_number(value * 100 if ratio else value, decimals, loc, sign)
    return number + (NBSP + "%" if loc == "fr" else "%")


def format_currency(value: float, currency: str = "EUR", decimals: int = 2,
                    locale: str = "fr") -> str:
    """1234.5 -> '1 234,50 €' (fr) / '€1,234.50' (en). Other currencies use the ISO code."""
    loc = _norm_locale(locale)
    symbols = {"EUR": "€", "USD": "$", "GBP": "£", "JPY": "¥", "CHF": "CHF"}
    sym = symbols.get(currency.upper(), currency.upper())
    number = format_number(abs(value), decimals, loc)
    neg = "-" if value < 0 else ""
    if loc == "fr":
        return f"{neg}{number}{NBSP}{sym}"
    if len(sym) > 1:
        return f"{neg}{sym}{NBSP}{number}"
    return f"{neg}{sym}{number}"


def _fold(text: str) -> str:
    """Lower case without accents ('Février' -> 'fevrier')."""
    import unicodedata

    return "".join(ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)).lower()


_MONTH_WORDS = {_fold(m): i + 1 for names in MONTHS.values() for i, m in enumerate(names)}
_MONTH_WORDS.update({a: i + 1 for i, abbrevs in enumerate((
    ("jan", "janv"), ("feb", "fev", "fevr"), ("mar",), ("apr", "avr"), (), ("jun",),
    ("jul", "juil"), ("aug", "aou"), ("sep", "sept"), ("oct",), ("nov",), ("dec",))) for a in abbrevs})
_IGNORED_WORDS = {"le", "the", "of", "lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche",
                  "monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday",
                  "mon", "tue", "tues", "wed", "thu", "thur", "thurs", "fri", "sat", "sun"}


def parse_date(value) -> _dt.date | None:
    """A date from a date or a string, or None (never raises). Reads ISO ('2025-10-14',
    '2025-10-14T09:30'), numeric day-first ('14/10/2025', '14.10.2025'; month-first only when
    unambiguous, '10/31/2025') and French / English month names ('14 octobre 2025',
    'mardi 1er juillet 2025', '14 oct. 2025', 'October 14, 2025', '14th October 2025').
    A month without a day ('octobre 2025') or free text ('T3 2025') gives None."""
    import re

    try:
        if isinstance(value, _dt.datetime):
            return value.date()
        if isinstance(value, _dt.date):
            return value
        if not isinstance(value, str) or not value.strip():
            return None
        text = value.strip()
        m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ].*)?", text)
        if m:
            return _dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.fullmatch(r"(\d{4})/(\d{1,2})/(\d{1,2})", text)
        if m:
            return _dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", text)
        if m:
            day, month, year = int(m[1]), int(m[2]), int(m[3])
            if month > 12 >= day:
                day, month = month, day
            return _dt.date(year, month, day)
        tokens = re.findall(r"\d+(?:er|re|st|nd|rd|th)?|[^\W\d_]+", _fold(text))
        tokens = [t for t in tokens if t not in _IGNORED_WORDS]
        kinds, values = [], []
        for t in tokens:
            if t in _MONTH_WORDS:
                kinds.append("m"); values.append(_MONTH_WORDS[t])
            elif re.fullmatch(r"\d{4}", t):
                kinds.append("y"); values.append(int(t))
            elif re.fullmatch(r"\d{1,2}(?:er|re|st|nd|rd|th)?", t):
                kinds.append("d"); values.append(int(re.match(r"\d+", t)[0]))
            else:
                return None
        if kinds == ["d", "m", "y"]:
            return _dt.date(values[2], values[1], values[0])
        if kinds == ["m", "d", "y"]:
            return _dt.date(values[2], values[0], values[1])
        return None
    except (ValueError, TypeError, OverflowError):
        return None


def date_label(value=None, locale: str = "fr", style: str = "long") -> str:
    """The text of a date for a document: None -> today; a date or a parseable string (see
    parse_date) -> format_date in the locale ('14 octobre 2025' / '14 October 2025'); any other
    text ('Automne 2025', 'T3 2025 - draft') -> as written."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return format_date(None, locale, style)
    d = parse_date(value)
    return format_date(d, locale, style) if d is not None else str(value).strip()


def format_date(d: _dt.date | None = None, locale: str = "fr", style: str = "long") -> str:
    """style: long ('6 octobre 2026' / '6 October 2026'), short ('06/10/2026' fr,
    '06/10/2026' en-GB style), month ('octobre 2026'), iso ('2026-10-06')."""
    loc = _norm_locale(locale)
    d = d or _dt.date.today()
    if style == "iso":
        return d.isoformat()
    if style == "short":
        return f"{d.day:02d}/{d.month:02d}/{d.year}"
    month = MONTHS[loc][d.month - 1]
    if style == "month":
        return f"{month} {d.year}"
    if loc == "fr":
        day = "1er" if d.day == 1 else str(d.day)
        return f"{day} {month} {d.year}"
    return f"{d.day} {month} {d.year}"


# Excel number formats are locale independent: Excel/LibreOffice display the user's own
# separators. Never store numbers as formatted text.
XLSX_FORMATS = {
    "integer": "#,##0",
    "number": "#,##0.00",
    "percent": "0.0%",
    "percent0": "0%",
    "eur": '#,##0.00 "€"',
    "eur0": '#,##0 "€"',
    "usd": '"$"#,##0.00',
    "keur": '#,##0,"k€"',         # thousands
    "meur": '#,##0.0,,"M€"',       # millions
    "date": "dd/mm/yyyy",
    "date_iso": "yyyy-mm-dd",
    "month": "mmm yyyy",
    "delta": "+#,##0.0;-#,##0.0;0.0",
}

LABELS = {
    "fr": {"page": "Page", "of": "sur", "prepared_by": "Préparé par", "date": "Date",
           "source": "Source", "contents": "Sommaire", "draft": "Brouillon"},
    "en": {"page": "Page", "of": "of", "prepared_by": "Prepared by", "date": "Date",
           "source": "Source", "contents": "Contents", "draft": "Draft"},
}


# --------------------------------------------------------------------------- python-docx
def _docx_set_font(style_or_run_font, name: str):
    """Set a font on a python-docx Font object for all scripts (ascii, hAnsi, eastAsia, cs)."""
    from docx.oxml.ns import qn
    style_or_run_font.name = name
    element = style_or_run_font.element  # rPr owner
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


class _FontProxy:
    def __init__(self, font, element):
        self._font, self.element = font, element

    @property
    def name(self):
        return self._font.name

    @name.setter
    def name(self, value):
        self._font.name = value


def docx_apply(document, locale: str = "fr"):
    """Apply Studio defaults to a python-docx Document: A4, margins, fonts, colours,
    heading sizes, list/table/caption styles, document language. Returns the document."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    loc = _norm_locale(locale)
    for section in document.sections:
        section.page_width, section.page_height = Cm(DOCX_PAGE["width_cm"]), Cm(DOCX_PAGE["height_cm"])
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(DOCX_PAGE["margin_cm"]))
        section.header_distance = section.footer_distance = Cm(1.0)

    styles = document.styles

    def style_font(style, size=None, bold=None, color=None, italic=None):
        _docx_set_font(_FontProxy(style.font, style.element), FONTS["office"])
        if size is not None:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if italic is not None:
            style.font.italic = italic
        if color is not None:
            style.font.color.rgb = RGBColor.from_string(PALETTE.get(color, color))

    normal = styles["Normal"]
    style_font(normal, DOCX_SIZES["body"], color="ink")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    # Language for spell-check and hyphenation.
    rpr = normal.element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = rpr.makeelement(qn("w:lang"), {})
        rpr.append(lang)
    lang.set(qn("w:val"), "fr-FR" if loc == "fr" else "en-GB")

    heading_spec = {1: ("h1", "navy", 18, 6), 2: ("h2", "navy", 12, 4), 3: ("h3", "blue", 10, 3)}
    for level, (key, color, before, after) in heading_spec.items():
        try:
            st = styles[f"Heading {level}"]
        except KeyError:
            st = styles.add_style(f"Heading {level}", WD_STYLE_TYPE.PARAGRAPH)
        style_font(st, DOCX_SIZES[key], bold=True, color=color, italic=False)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name, key, color, bold in (("Title", "title", "navy", True),
                                   ("Subtitle", "subtitle", "slate", False)):
        try:
            st = styles[name]
        except KeyError:
            continue
        style_font(st, DOCX_SIZES[key], bold=bold, color=color, italic=False)
        # The default template draws a border under Title: remove it.
        ppr = st.element.find(qn("w:pPr"))
        if ppr is not None:
            for bdr in ppr.findall(qn("w:pBdr")):
                ppr.remove(bdr)
    for name in ("Caption",):
        try:
            style_font(styles[name], DOCX_SIZES["caption"], bold=False, italic=True, color="slate")
        except KeyError:
            pass
    for name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
        try:
            style_font(styles[name], DOCX_SIZES["body"], color="ink")
        except KeyError:
            pass
    return document


def docx_title_block(document, title: str, subtitle: str | None = None, author: str | None = None,
                     date: _dt.date | str | None = None, locale: str = "fr", page_break: bool = True):
    """Title page: title, subtitle, a navy rule, then 'Prepared by' and date lines. date: a date,
    any text (dates are written in the locale's long form, other text as is; see date_label) or
    None (today)."""
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    loc = _norm_locale(locale)
    lab = LABELS[loc]
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(120)
    t = document.add_paragraph(title, style="Title")
    if subtitle:
        document.add_paragraph(subtitle, style="Subtitle")
    rule = document.add_paragraph()
    ppr = rule._p.get_or_add_pPr()
    pbdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = pbdr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "12",
                                                qn("w:space"): "1", qn("w:color"): PALETTE["navy"]})
    pbdr.append(bottom)
    ppr.append(pbdr)
    meta = []
    if author:
        meta.append(f"{lab['prepared_by']} : {author}" if loc == "fr" else f"{lab['prepared_by']}: {author}")
    meta.append(date_label(date, loc, "long"))
    for line in meta:
        p = document.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(DOCX_SIZES["body"])
        r.font.color.rgb = RGBColor.from_string(PALETTE["slate"])
    if page_break:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return t


def _docx_field(paragraph, instr: str, placeholder: str = "1"):
    """Append a complex field (PAGE, NUMPAGES, DATE...) to a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def run_with(child):
        r = paragraph.add_run()
        r._r.append(child)
        return r

    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = f" {instr} "
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    runs = [run_with(b), run_with(i), run_with(s)]
    runs.append(paragraph.add_run(placeholder))
    runs.append(run_with(e))
    return runs


def _docx_clear_style_tabs(paragraph):
    """Cancel the tab stops inherited from the paragraph style (Header/Footer styles define
    centre and right tabs that would otherwise capture our tab characters)."""
    from docx.enum.text import WD_TAB_ALIGNMENT

    style = paragraph.style
    while style is not None:
        for ts in style.paragraph_format.tab_stops:
            paragraph.paragraph_format.tab_stops.add_tab_stop(ts.position, WD_TAB_ALIGNMENT.CLEAR)
        style = style.base_style


def docx_footer(document, left_text: str = "", locale: str = "fr", header_text: str | None = None):
    """Footer: ONE right-aligned paragraph '<left_text> · Page X sur Y' / '... · Page X of Y',
    grey 8 pt, thin rule above. Whatever the footer held before is removed first (no second
    page number, no number glued to the title). Optional header text right-aligned.
    Applies to every section."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    loc = _norm_locale(locale)
    lab = LABELS[loc]
    prefix = f"{left_text.strip()} · " if left_text and left_text.strip() else ""

    def clear(part):
        for child in list(part._element):
            part._element.remove(child)
        return part.add_paragraph()

    for section in document.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = clear(footer)
        _docx_clear_style_tabs(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ppr = p._p.get_or_add_pPr()
        pbdr = ppr.makeelement(qn("w:pBdr"), {})
        pbdr.append(pbdr.makeelement(qn("w:top"), {qn("w:val"): "single", qn("w:sz"): "4",
                                                   qn("w:space"): "4", qn("w:color"): PALETTE["rule"]}))
        ppr.append(pbdr)
        runs = [p.add_run(prefix + lab["page"] + " ")]
        runs += _docx_field(p, "PAGE")
        runs.append(p.add_run(f" {lab['of']} "))
        runs += _docx_field(p, "NUMPAGES")
        for r in runs:
            r.font.size = Pt(DOCX_SIZES["footer"])
            r.font.color.rgb = RGBColor.from_string(PALETTE["grey"])
            r.font.name = FONTS["office"]
        if header_text:
            header = section.header
            header.is_linked_to_previous = False
            hp = clear(header)
            _docx_clear_style_tabs(hp)
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hr = hp.add_run(header_text)
            hr.font.size = Pt(DOCX_SIZES["footer"])
            hr.font.color.rgb = RGBColor.from_string(PALETTE["grey"])
    return document


def docx_style_table(table, header: bool = True, zebra: bool = True, font_size: float | None = None,
                     numeric_cols: list[int] | None = None):
    """Header row navy with white bold text, light rule borders, optional zebra rows,
    numeric columns right-aligned. Marks the header row as repeating on each page."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    size = Pt(font_size or DOCX_SIZES["table"])
    numeric_cols = numeric_cols or []
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if edge not in ("left", "right", "insideV") else "nil")
        el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), PALETTE["rule"])
        borders.append(el)
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    tblpr.append(borders)

    def shade(cell, color):
        tcpr = cell._tc.get_or_add_tcPr()
        for old in tcpr.findall(qn("w:shd")):
            tcpr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        tcpr.append(shd)

    for r_idx, row in enumerate(table.rows):
        is_header = header and r_idx == 0
        if is_header:
            trpr = row._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
            trpr.append(th)
        for c_idx, cell in enumerate(row.cells):
            if is_header:
                shade(cell, PALETTE["navy"])
            elif zebra and r_idx % 2 == 0:
                shade(cell, PALETTE["mist"])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                if c_idx in numeric_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs:
                    run.font.size = size
                    run.font.name = FONTS["office"]
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(PALETTE["white"])
    return table


# --------------------------------------------------------------------------- python-pptx
def pptx_new(widescreen: bool = True):
    """New python-pptx Presentation, 16:9, with the theme fonts set to Arial and the theme
    colour scheme replaced by the Studio palette (so built-in layouts, tables and charts
    pick it up)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(PPTX_SLIDE["width_in"]), Inches(PPTX_SLIDE["height_in"])
    pptx_apply_theme(prs)
    return prs


def pptx_apply_theme(prs):
    """Rewrite the theme of the first slide master: fonts -> Arial, colours -> palette.
    Safe on existing decks too (only the theme part is changed)."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    master_part = prs.slide_masters[0].part
    theme_part = master_part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    scheme = {"dk1": PALETTE["ink"], "lt1": PALETTE["white"], "dk2": PALETTE["navy"],
              "lt2": PALETTE["mist"], "accent1": CHART_COLORS[0], "accent2": CHART_COLORS[1],
              "accent3": CHART_COLORS[2], "accent4": CHART_COLORS[3], "accent5": CHART_COLORS[4],
              "accent6": CHART_COLORS[5], "hlink": PALETTE["blue"], "folHlink": "7D5BA6"}
    clr = root.find(".//a:clrScheme", ns)
    if clr is not None:
        clr.set("name", "Studio")
        for key, value in scheme.items():
            el = clr.find(f"a:{key}", ns)
            if el is None:
                continue
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{{{ns['a']}}}srgbClr", val=value)
    font_scheme = root.find(".//a:fontScheme", ns)
    if font_scheme is not None:
        font_scheme.set("name", "Studio")
        for which in ("a:majorFont", "a:minorFont"):
            latin = font_scheme.find(f"{which}/a:latin", ns)
            if latin is not None:
                latin.set("typeface", FONTS["office"])
                for attr in ("panose", "pitchFamily", "charset"):
                    latin.attrib.pop(attr, None)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return prs


def pptx_style_text(text_frame, size: float | None = None, color: str = "ink", bold: bool | None = None):
    """Set font family/size/colour on every run of a python-pptx text frame."""
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    for p in text_frame.paragraphs:
        for r in p.runs:
            r.font.name = FONTS["office"]
            if size:
                r.font.size = Pt(size)
            r.font.color.rgb = RGBColor.from_string(PALETTE.get(color, color))
            if bold is not None:
                r.font.bold = bold


def pptx_footer(slide, prs, text: str = "", number: int | None = None):
    """Small grey footer text bottom-left and slide number bottom-right (as text boxes;
    python-pptx does not activate the layout's footer placeholders)."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    m = Inches(PPTX_SLIDE["margin_in"])
    y = prs.slide_height - Inches(0.45)
    h = Inches(0.3)
    boxes = []
    if text:
        tb = slide.shapes.add_textbox(m, y, prs.slide_width // 2, h)
        boxes.append((tb, text, PP_ALIGN.LEFT))
    if number is not None:
        tb = slide.shapes.add_textbox(prs.slide_width - m - Inches(1.0), y, Inches(1.0), h)
        boxes.append((tb, str(number), PP_ALIGN.RIGHT))
    for tb, value, align in boxes:
        tb.name = "Studio footer"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = value
        r.font.size = Pt(PPTX_SIZES["footer"])
        r.font.name = FONTS["office"]
        r.font.color.rgb = RGBColor.from_string(PALETTE["grey"])
    return [b[0] for b in boxes]


def pptx_style_table(table, header: bool = True, zebra: bool = True, font_size: float = 12,
                     numeric_cols: list[int] | None = None):
    """Style a python-pptx table (graphic_frame.table)."""
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    numeric_cols = numeric_cols or []
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            is_header = header and r_idx == 0
            cell.fill.solid()
            if is_header:
                cell.fill.fore_color.rgb = RGBColor.from_string(PALETTE["navy"])
            elif zebra and r_idx % 2 == 0:
                cell.fill.fore_color.rgb = RGBColor.from_string(PALETTE["mist"])
            else:
                cell.fill.fore_color.rgb = RGBColor.from_string(PALETTE["white"])
            for p in cell.text_frame.paragraphs:
                if c_idx in numeric_cols:
                    p.alignment = PP_ALIGN.RIGHT
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = FONTS["office"]
                    run.font.bold = bool(is_header)
                    run.font.color.rgb = RGBColor.from_string(PALETTE["white" if is_header else "ink"])
    return table


# --------------------------------------------------------------------------- openpyxl
def xlsx_header(ws, row: int = 1, first_col: int = 1, last_col: int | None = None,
                freeze: bool = True):
    """Style a header row (navy fill, white bold Arial, bottom border), freeze panes below
    it and turn on the autofilter over the used range."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    last_col = last_col or ws.max_column
    fill = PatternFill("solid", fgColor=PALETTE["navy"])
    font = Font(name=FONTS["office"], bold=True, color=PALETTE["white"], size=XLSX_SIZES["header"])
    border = Border(bottom=Side(style="thin", color=PALETTE["rule"]))
    for col in range(first_col, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.fill, c.font, c.border = fill, font, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=first_col)
    ws.auto_filter.ref = f"{get_column_letter(first_col)}{row}:{get_column_letter(last_col)}{max(ws.max_row, row)}"
    return ws


def xlsx_body_font(ws, min_row: int = 2):
    """Arial 10 ink on every non-empty cell from min_row (keeps bold/italic)."""
    from openpyxl.styles import Font

    for row in ws.iter_rows(min_row=min_row):
        for c in row:
            if c.value is not None:
                f = c.font
                c.font = Font(name=FONTS["office"], size=XLSX_SIZES["body"], bold=f.bold,
                              italic=f.italic, color=PALETTE["ink"])
    return ws


def xlsx_autofit(ws, min_width: float = 8, max_width: float = 60):
    """Approximate column widths from the longest displayed value per column."""
    from openpyxl.utils import get_column_letter

    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                length = 12
            elif isinstance(v, float):
                length = len(f"{v:,.2f}")
            elif isinstance(v, (_dt.date, _dt.datetime)):
                length = 11
            else:
                length = max(len(line) for line in str(v).split("\n"))
            widths[c.column] = max(widths.get(c.column, 0), length)
    for col, length in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, length * 1.15 + 2))
    return ws


def xlsxwriter_formats(workbook) -> dict:
    """Common xlsxwriter formats: header, text, integer, number, percent, eur, date, title."""
    base = {"font_name": FONTS["office"], "font_size": XLSX_SIZES["body"], "font_color": hex_("ink")}
    fmts = {
        "title": workbook.add_format({**base, "bold": True, "font_size": XLSX_SIZES["title"], "font_color": hex_("navy")}),
        "header": workbook.add_format({**base, "bold": True, "font_color": "#FFFFFF", "bg_color": hex_("navy"),
                                       "align": "center", "valign": "vcenter", "text_wrap": True,
                                       "bottom": 1, "bottom_color": hex_("rule")}),
        "text": workbook.add_format(base),
    }
    for key in ("integer", "number", "percent", "eur", "date"):
        fmts[key] = workbook.add_format({**base, "num_format": XLSX_FORMATS[key]})
    return fmts


# --------------------------------------------------------------------------- reportlab
def reportlab_fonts() -> str:
    """Register Liberation Sans (regular/bold/italic/bold-italic) as family 'StudioSans' in
    reportlab and return the family name. Unlike the built-in Helvetica it covers all
    Unicode Latin text (accents, €, no-break spaces, Greek, Cyrillic)."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name = FONTS["reportlab"]
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    faces = {"": "regular", "-Bold": "bold", "-Italic": "italic", "-BoldItalic": "bold_italic"}
    for suffix, key in faces.items():
        pdfmetrics.registerFont(TTFont(name + suffix, FONT_FILES[key]))
    addMapping(name, 0, 0, name)
    addMapping(name, 1, 0, name + "-Bold")
    addMapping(name, 0, 1, name + "-Italic")
    addMapping(name, 1, 1, name + "-BoldItalic")
    return name


# --------------------------------------------------------------------------- matplotlib
def mpl_apply():
    """Set matplotlib rcParams to the Studio chart style (call before creating figures)."""
    import matplotlib
    from cycler import cycler

    matplotlib.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": FONTS["render_fallback"],
        "font.size": 10,
        "axes.titlesize": 13,
        "axes.titleweight": "bold",
        "axes.titlelocation": "left",
        "axes.titlecolor": hex_("ink"),
        "axes.labelsize": 10,
        "axes.labelcolor": hex_("slate"),
        "axes.edgecolor": hex_("grey"),
        "axes.linewidth": 0.8,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.grid": True,
        "axes.grid.axis": "y",
        "axes.axisbelow": True,
        "axes.prop_cycle": cycler(color=[hex_(c) for c in CHART_COLORS]),
        "grid.color": hex_("rule"),
        "grid.linewidth": 0.6,
        "xtick.color": hex_("slate"),
        "ytick.color": hex_("slate"),
        "xtick.labelsize": 9,
        "ytick.labelsize": 9,
        "legend.frameon": False,
        "legend.fontsize": 9,
        "figure.facecolor": "white",
        "figure.dpi": 100,
        "savefig.dpi": DPI,
        "savefig.bbox": "tight",
        "savefig.facecolor": "white",
        "lines.linewidth": 2.0,
        "svg.fonttype": "none",
        "axes.formatter.use_locale": False,
    })
    return matplotlib.rcParams


def mpl_number_formatter(kind: str = "number", decimals: int = 0, locale: str = "fr"):
    """matplotlib tick formatter: kind in number|percent (ratio)|eur."""
    from matplotlib.ticker import FuncFormatter

    def fmt(x, _pos=None):
        if kind == "percent":
            return format_percent(x, decimals, locale)
        if kind == "eur":
            return format_currency(x, "EUR", decimals, locale)
        return format_number(x, decimals, locale)

    return FuncFormatter(fmt)


# --------------------------------------------------------------------------- CLI
def _swatch(path: str):
    from PIL import Image, ImageDraw, ImageFont

    items = list(PALETTE.items())
    w, h = 180, 90
    img = Image.new("RGB", (w * 7, h * ((len(items) + 6) // 7)), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(FONT_FILES["regular"], 14)
    except OSError:
        font = ImageFont.load_default()
    for i, (name, value) in enumerate(items):
        x, y = (i % 7) * w, (i // 7) * h
        draw.rectangle([x + 4, y + 4, x + w - 4, y + h - 30], fill=hex_(value), outline=hex_("rule"))
        draw.text((x + 6, y + h - 26), f"{name} #{value}", fill=hex_("ink"), font=font)
    img.save(path)
    return path


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json", action="store_true", help="print brand constants as JSON")
    ap.add_argument("--swatch", metavar="PNG", help="write a palette swatch image")
    ap.add_argument("--format", type=float, metavar="NUMBER", help="format a number")
    ap.add_argument("--locale", default="fr", choices=["fr", "en"])
    ap.add_argument("--decimals", type=int, default=2)
    args = ap.parse_args(argv)
    if args.swatch:
        print(json.dumps({"swatch": _swatch(args.swatch)}))
    if args.format is not None:
        print(json.dumps({"number": format_number(args.format, args.decimals, args.locale),
                          "percent": format_percent(args.format, args.decimals, args.locale, ratio=False),
                          "currency": format_currency(args.format, "EUR", args.decimals, args.locale),
                          "date": format_date(None, args.locale)}, ensure_ascii=False))
    if args.json or (not args.swatch and args.format is None):
        print(json.dumps({"palette": {k: "#" + v for k, v in PALETTE.items()},
                          "chart_colors": ["#" + c for c in CHART_COLORS],
                          "sequential": ["#" + c for c in SEQUENTIAL],
                          "diverging": ["#" + c for c in DIVERGING],
                          "fonts": FONTS, "docx_sizes_pt": DOCX_SIZES, "pptx_sizes_pt": PPTX_SIZES,
                          "xlsx_formats": XLSX_FORMATS, "figure_sizes_in": FIGURE_SIZES, "dpi": DPI},
                         indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - bad input: one JSON error on stdout, no traceback
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
