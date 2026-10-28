#!/usr/bin/env python3
"""Recalculate every formula of an .xlsx with LibreOffice headless, store the computed
values in the ORIGINAL file (formulas, styles, charts, everything else untouched) and
report formula errors (#REF!, #DIV/0!, #VALUE!, #NAME?, #N/A, #NUM!, #NULL!) as JSON.

Why: openpyxl / xlsxwriter write formulas without results. Until something calculates
them, previews, pandas, `openpyxl(data_only=True)` and many viewers show empty cells.

Usage:
  python3 recalc.py BOOK.xlsx                    # recalc in place + report
  python3 recalc.py BOOK.xlsx --out RESULT.xlsx  # write the recalculated copy elsewhere
  python3 recalc.py BOOK.xlsx --check-only       # report only, file unchanged
  python3 recalc.py BOOK.xlsx --max-errors 200

Report (the checker contract of the xlsx skill): {"ok", "file", "checks": [{"name", "ok", "detail"}],
"next_step", "formulas", "errors_total", "errors_by_type", "errors": [{sheet, cell, error,
formula}], "unsupported_functions": [...], "warnings": [...]}
Checks: opens (LibreOffice opened and recalculated it), formula_errors (no #DIV/0!, #REF!,
#NAME?...), live_formulas (rows labelled Total/Sum/Average/Moyenne hold formulas, not typed
numbers), numbers_as_text (no number stored as text inside a range a formula uses: SUM and
friends count it as 0).
Exit code: 0 = every check passed, 2 = a check failed (fix it and re-run), 1 = usage error
(no such file, not .xlsx/.xlsm, bad arguments).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!")
XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


CMD = "python3 /skills/xlsx/scripts/recalc.py"
TOTAL_LABEL = re.compile(r"^\s*(grand total|sub-?total|sous-total|total g[ée]n[ée]ral|totals?|sum|somme|average|"
                         r"moyenne|mean)\b", re.I)
NUMERIC_TEXT = re.compile(r"^\s*[-+(]?\s*(\d[\d\s\u00a0\u202f]*([.,]\d+)*|[.,]\d+)\s*\)?\s*(%|€|\$|£|EUR|USD|k€|M€)?\s*$")
REF = re.compile(r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][\w.]*))!)?(?<![A-Za-z0-9_.])\$?([A-Z]{1,3})\$?(\d+)"
                 r"(?::\$?([A-Z]{1,3})\$?(\d+))?(?![\w(])")
MAX_CELLS = 300000


class Parser(argparse.ArgumentParser):
    def error(self, message):  # usage error: one JSON object, exit 1 (argparse would exit 2)
        print(json.dumps({"ok": False, "error": message, "usage": self.format_usage().strip()}))
        sys.exit(1)


def content_checks(path):
    """live_formulas and numbers_as_text checks on the formulas of the file (openpyxl)."""
    import openpyxl
    from openpyxl.utils.cell import column_index_from_string

    wb = openpyxl.load_workbook(path)
    hard, as_text = [], []
    seen = 0
    for ws in wb.worksheets:
        if seen > MAX_CELLS:
            break
        numbers = {}            # col -> set(rows) of numeric constants
        texts = {}              # (row, col) -> value of numeric-looking strings
        formulas = []
        rows = list(ws.iter_rows())
        for row in rows:
            for c in row:
                seen += 1
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    formulas.append(v)
                elif isinstance(v, bool) or v is None:
                    continue
                elif isinstance(v, (int, float)):
                    numbers.setdefault(c.column, set()).add(c.row)
                elif isinstance(v, str) and NUMERIC_TEXT.match(v) and any(ch.isdigit() for ch in v):
                    texts[(c.row, c.column)] = v
        for row in rows:
            label = next((c for c in row[:3] if isinstance(c.value, str) and not c.value.startswith("=")), None)
            if label is None or not TOTAL_LABEL.match(label.value):
                continue
            for c in row:
                if c.column <= label.column or isinstance(c.value, bool) or not isinstance(c.value, (int, float)):
                    continue
                above = [r for r in numbers.get(c.column, ()) if r < c.row]
                if len(above) >= 2:
                    hard.append(f"{ws.title}!{c.coordinate} = {c.value} ({label.value.strip()[:20]!r} row)")
        if texts:
            for f in formulas:
                for m in REF.finditer(f):
                    sheet = (m.group(1) or "").replace("''", "'") or m.group(2)
                    if sheet and sheet != ws.title:
                        continue
                    c1, r1 = column_index_from_string(m.group(3)), int(m.group(4))
                    c2, r2 = (column_index_from_string(m.group(5)), int(m.group(6))) if m.group(5) else (c1, r1)
                    for (r, c), v in texts.items():
                        if min(r1, r2) <= r <= max(r1, r2) and min(c1, c2) <= c <= max(c1, c2):
                            cell = f"{ws.title}!{openpyxl.utils.get_column_letter(c)}{r}"
                            entry = f"{cell} = {v!r}"
                            if entry not in as_text:
                                as_text.append(entry)
    return hard, as_text


def build_checks(opened, error, total, by_type, errors, hard, as_text):
    if not opened:
        return [{"name": "opens", "ok": False, "detail": error}]
    checks = [{"name": "opens", "ok": True, "detail": "LibreOffice opened and recalculated the workbook"}]
    if total:
        first = "; ".join(f"{e['sheet']}!{e['cell']} {e['error']} {e.get('formula') or ''}".strip() for e in errors[:5])
        checks.append({"name": "formula_errors", "ok": False,
                       "detail": f"{total} formula errors {by_type}: {first}. Fix the range / wrap with IFERROR / "
                                 "replace functions LibreOffice does not know"})
    else:
        checks.append({"name": "formula_errors", "ok": True, "detail": "no formula error"})
    checks.append({"name": "live_formulas", "ok": not hard,
                   "detail": ("typed numbers in total rows (the sheet looks computed, write formulas such as "
                              "=SUM(B2:B9) instead): " + "; ".join(hard[:6])) if hard else "totals are formulas"})
    checks.append({"name": "numbers_as_text", "ok": not as_text,
                   "detail": ("numbers stored as text inside ranges that formulas use (they count as 0): "
                              + "; ".join(as_text[:6]) + ". Write them as numbers (openpyxl: int/float values; "
                              "xlsxwriter: write_number)") if as_text else "no number stored as text in a formula range"})
    return checks


def next_step(path, checks):
    failed = [c["name"] for c in checks if not c["ok"]]
    if not failed:
        return ("OK: all checks passed. Check a few key values (inspect_xlsx.py --values) and, for a formatted report, "
                f"call view_file on {path} before you answer.")
    return (f"NOT DONE: fix the failing checks ({', '.join(failed)}), save, then run `{CMD} {path}` again (at most 3 "
            "rounds). Do not say the workbook is ready while ok is false.")


def q(tag):
    return "{%s}%s" % (S, tag)


def is_error(v):
    return isinstance(v, str) and (v in ERRORS or re.match(r"^Err:\d{3}$", v) is not None)


def lo_recalc(src, workdir, timeout):
    profile = os.path.join(workdir, "profile")
    os.makedirs(os.path.join(profile, "user"), exist_ok=True)
    with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as fh:
        fh.write(XCU)
    outdir = os.path.join(workdir, "out")
    os.makedirs(outdir)
    # copy to a neutral name so odd characters in the user's filename cannot break soffice
    tmp_in = os.path.join(workdir, "in" + os.path.splitext(src)[1].lower())
    shutil.copyfile(src, tmp_in)
    cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--nolockcheck",
           "--nodefault", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir, tmp_in]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"LibreOffice timed out after {timeout}s") from exc
    out = os.path.join(outdir, "in.xlsx")
    if not os.path.isfile(out):
        raise RuntimeError(f"LibreOffice produced no output: {res.stdout} {res.stderr}")
    return out


def sheet_parts(zf):
    """[(sheet name, part name)] in workbook order."""
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    out = []
    for sh in wb.find(q("sheets")):
        target = targets.get(sh.get("{%s}id" % R), "")
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        out.append((sh.get("name"), part))
    return out


def expand_range(ref):
    from openpyxl.utils.cell import range_boundaries, get_column_letter

    min_col, min_row, max_col, max_row = range_boundaries(ref)
    return [f"{get_column_letter(c)}{r}" for r in range(min_row, max_row + 1) for c in range(min_col, max_col + 1)]


def to_cell_xml(c, value):
    """Store `value` as the cached result of formula cell `c` (lxml element)."""
    from openpyxl.utils.datetime import to_excel

    for v in c.findall(q("v")):
        c.remove(v)
    for isel in c.findall(q("is")):
        c.remove(isel)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        value = to_excel(value)
    if isinstance(value, dt.timedelta):
        value = value.total_seconds() / 86400
    v = etree.Element(q("v"))
    f = c.find(q("f"))
    if f is not None:
        f.addnext(v)  # schema order: f, v, is, extLst
    else:
        c.insert(0, v)
    if value is None:
        c.set("t", "str")
        v.text = ""
    elif isinstance(value, bool):
        c.set("t", "b")
        v.text = "1" if value else "0"
    elif isinstance(value, (int, float)):
        c.attrib.pop("t", None)
        v.text = repr(int(value)) if isinstance(value, int) or float(value).is_integer() and abs(value) < 1e15 else repr(float(value))
    elif is_error(value):
        c.set("t", "e")
        v.text = value if value in ERRORS else "#VALUE!"
    else:
        c.set("t", "str")
        v.text = str(value)


def main(argv=None) -> int:
    ap = Parser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--out")
    ap.add_argument("--check-only", action="store_true")
    ap.add_argument("--max-errors", type=int, default=100)
    ap.add_argument("--timeout", type=int, default=180)
    args = ap.parse_args(argv)

    import openpyxl

    src = os.path.abspath(args.file)
    if not os.path.isfile(src):
        print(json.dumps({"ok": False, "file": src, "error": f"file not found: {src}"}))
        return 1
    if not src.lower().endswith((".xlsx", ".xlsm")):
        print(json.dumps({"ok": False, "file": src, "error": "recalc.py handles .xlsx / .xlsm only"}))
        return 1
    work = tempfile.mkdtemp(prefix="recalc.", dir="/tmp")
    try:
        try:
            zipfile.ZipFile(src).testzip()
            computed_path = lo_recalc(src, work, args.timeout)
            values = openpyxl.load_workbook(computed_path, data_only=True)
            lo_formulas = openpyxl.load_workbook(computed_path, data_only=False)
        except Exception as exc:  # noqa: BLE001
            checks = build_checks(False, f"cannot open/recalculate the workbook: {exc}", 0, {}, [], [], [])
            print(json.dumps({"ok": False, "file": src, "checks": checks, "next_step": next_step(src, checks),
                              "error": str(exc)}))
            return 2

        zin = zipfile.ZipFile(src)
        parts = sheet_parts(zin)
        new_parts = {}
        errors, by_type, formulas = [], {}, 0
        unsupported = set()
        warnings = []
        for name, part in parts:
            if part not in zin.namelist():
                continue
            root = etree.fromstring(zin.read(part))
            if name not in values.sheetnames:
                warnings.append(f"sheet {name!r} missing from the LibreOffice result")
                continue
            ws_v, ws_f = values[name], lo_formulas[name]
            data = root.find(q("sheetData"))
            if data is None:
                continue
            array_cells = set()
            for c in data.iter(q("c")):
                f = c.find(q("f"))
                if f is not None and f.get("t") == "array" and f.get("ref"):
                    array_cells.update(expand_range(f.get("ref")))
            for c in data.iter(q("c")):
                ref = c.get("r")
                f = c.find(q("f"))
                if f is None and ref not in array_cells:
                    continue
                if f is not None:
                    formulas += 1
                value = ws_v[ref].value
                if is_error(value):
                    by_type[value] = by_type.get(value, 0) + 1
                    if len(errors) < args.max_errors:
                        errors.append({"sheet": name, "cell": ref, "error": value,
                                       "formula": ("=" + f.text) if f is not None and f.text else None})
                    lo_f = ws_f[ref].value
                    if value == "#NAME?" and isinstance(lo_f, str):
                        for fn in re.findall(r"([a-z][a-z0-9.]*)\(", lo_f):
                            unsupported.add(fn.upper())
                to_cell_xml(c, value)
            new_parts[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

        # make Excel recalculate on open as well
        wb_root = etree.fromstring(zin.read("xl/workbook.xml"))
        calc = wb_root.find(q("calcPr"))
        if calc is None:
            calc = etree.SubElement(wb_root, q("calcPr"))
        calc.set("fullCalcOnLoad", "1")
        new_parts["xl/workbook.xml"] = etree.tostring(wb_root, xml_declaration=True, encoding="UTF-8", standalone=True)

        out = None
        if not args.check_only:
            out = os.path.abspath(args.out) if args.out else src
            fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(out))
            os.close(fd)
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    zout.writestr(item, new_parts.get(item.filename, zin.read(item.filename)))
            zin.close()
            shutil.move(tmp, out)
        if unsupported:
            warnings.append("functions unknown to LibreOffice 7.4 (#NAME?): " + ", ".join(sorted(unsupported))
                            + " - Excel may compute them, but prefer classic equivalents (INDEX/MATCH, SUMIFS, IFERROR)")
        total = sum(by_type.values())
        try:
            hard, as_text = content_checks(src)
        except Exception as exc:  # noqa: BLE001
            hard, as_text = [], []
            warnings.append(f"content checks skipped: {exc}")
        checks = build_checks(True, None, total, by_type, errors, hard, as_text)
        ok = all(c["ok"] for c in checks)
        path = out or src
        report = {"ok": ok, "file": path, "checks": checks, "next_step": next_step(path, checks),
                  "recalculated": not args.check_only, "formulas": formulas,
                  "errors_total": total, "errors_by_type": by_type, "errors": errors,
                  "unsupported_functions": sorted(unsupported), "warnings": warnings}
        print(json.dumps(report, indent=2, ensure_ascii=False))
        return 0 if ok else 2
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - bad input: one JSON error on stdout, no traceback
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
